"""Validate an uploaded resume file and extract its text.

Supported: PDF (text-based) and DOCX. Scanned/image-only PDFs have no text
layer and are rejected with a clear message (no OCR is performed).
"""

import io
import re
import zipfile

# Vercel serverless functions reject request bodies larger than 4.5 MB, so the
# limit must stay below that (multipart adds a little overhead).
MAX_FILE_BYTES = 4 * 1024 * 1024        # 4 MB
MAX_PDF_PAGES = 15
MAX_TEXT_CHARS = 60_000
MAX_DOCX_UNCOMPRESSED = 50 * 1024 * 1024  # zip-bomb guard
MIN_TEXT_CHARS = 80

ALLOWED_EXTENSIONS = {".pdf": "pdf", ".docx": "docx"}


class ResumeFileError(ValueError):
    """The uploaded file is invalid or unreadable. Message is safe to show."""


def detect_file_type(filename, data):
    name = (filename or "").lower()
    ext = next((e for e in ALLOWED_EXTENSIONS if name.endswith(e)), None)

    if ext is None:
        raise ResumeFileError("Unsupported file type. Upload a PDF or DOCX file.")

    if not data:
        raise ResumeFileError("The uploaded file is empty.")

    if len(data) > MAX_FILE_BYTES:
        raise ResumeFileError("The file is too large. The maximum size is 4 MB.")

    # Verify the content really matches the extension.
    if ext == ".pdf" and not data.lstrip()[:5].startswith(b"%PDF"):
        raise ResumeFileError("This file is not a valid PDF.")

    if ext == ".docx":
        if not data.startswith(b"PK\x03\x04"):
            raise ResumeFileError("This file is not a valid DOCX document.")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                names = archive.namelist()
                total = sum(info.file_size for info in archive.infolist())
        except zipfile.BadZipFile as exc:
            raise ResumeFileError("This file is not a valid DOCX document.") from exc
        if "word/document.xml" not in names:
            raise ResumeFileError("This file is not a valid DOCX document.")
        if total > MAX_DOCX_UNCOMPRESSED:
            raise ResumeFileError("This DOCX file is too large to process.")

    return ALLOWED_EXTENSIONS[ext]


def _clean(text):
    text = text.replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()[:MAX_TEXT_CHARS]


def _extract_pdf(data):
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ResumeFileError("Password-protected PDFs are not supported.")
        pages = reader.pages[:MAX_PDF_PAGES]
        return "\n".join((page.extract_text() or "") for page in pages)
    except ResumeFileError:
        raise
    except (PyPdfError, ValueError, KeyError, OSError, RecursionError) as exc:
        raise ResumeFileError("The PDF could not be read. It may be corrupted.") from exc


def _extract_docx(data):
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several unrelated types
        raise ResumeFileError("The DOCX file could not be read.") from exc

    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def extract_resume_text(filename, data):
    """Return (file_type, cleaned_text). Raises ResumeFileError on bad input."""

    file_type = detect_file_type(filename, data)
    raw = _extract_pdf(data) if file_type == "pdf" else _extract_docx(data)
    text = _clean(raw)

    if len(text) < MIN_TEXT_CHARS:
        raise ResumeFileError(
            "No readable text was found. If this is a scanned or image-only "
            "resume, export a text-based PDF or upload a DOCX instead."
        )

    return file_type, text
